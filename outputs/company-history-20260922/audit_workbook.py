import json, pathlib, re, zipfile
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from PIL import Image, ImageDraw, ImageFont

out=pathlib.Path(__file__).parent
path=out/'가스펠플러스_사업실적_대표경력_검토용.xlsx'
w=load_workbook(path)
original={s.title:list(s.values) for s in w if s!=w.active}
fixes={'C022':'2025~종료 미확인','P017':'2019 (회사자료 2016, 확인 필요)','P014':'2020 (회사자료 2021, 확인 필요)'}
changed=[]
for row in w.active.iter_rows(min_row=5):
    ident=row[1].comment.text.split(' | ')[0]
    if ident in fixes:
        row[3].value=fixes[ident]
        row[3].fill=PatternFill('solid',fgColor='FFF2D5')
        changed.append(row[3].coordinate)
w.save(path)

# Validate the nine simple summary formulas and restore caches lost on save.
cache={}
for row in w['검토안내']:
    for c in row:
        if c.data_type!='f':continue
        m=re.fullmatch(r'=COUNTA\(\'([^\']+)\'!(\w+:\w+)\)',c.value)
        if m:
            sheet,area=m.groups();value=sum(x.value is not None for rr in w[sheet][area] for x in rr)
        else:
            m=re.fullmatch(r'=COUNTIF\(\'([^\']+)\'!(\w+:\w+),"([^"]+)"\)',c.value)
            assert m,c.value
            sheet,area,criterion=m.groups();value=sum(x.value==criterion for rr in w[sheet][area] for x in rr)
        cache[c.coordinate]=value
with zipfile.ZipFile(path) as z:parts={i.filename:z.read(i.filename) for i in z.infolist()}
ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
sheetfile=f'xl/worksheets/sheet{w.sheetnames.index("검토안내")+1}.xml'
root=ET.fromstring(parts[sheetfile])
for c in root.findall('.//s:c',ns):
    if c.attrib['r'] in cache:c.find('s:v',ns).text=str(cache[c.attrib['r']])
parts[sheetfile]=ET.tostring(root,encoding='utf-8')
with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
    for name,data in parts.items():z.writestr(name,data)
with zipfile.ZipFile(path) as z:
    assert z.testzip() is None
    for name in z.namelist():
        if name.endswith(('.xml','.rels')):ET.fromstring(z.read(name))
v=load_workbook(path)
assert all(s.freeze_panes is None and s.auto_filter.ref is None and not s.tables for s in v)
assert [s.title for s in v if s.sheet_state=='visible']==['참여사업_요약']
assert v.active.max_row==57 and v.active.max_column==4
assert all(list(v[name].values)==rows for name,rows in original.items())
assert not any(c.data_type=='e' for s in v for row in s for c in row)
ids={i for row in v.active.iter_rows(min_row=5) for i in re.findall(r'\b[CP]\d{3}\b',row[1].comment.text)}
expected={r[0] for n in ['회사실적_문서근거','대표경력_개인'] for r in v[n].iter_rows(min_row=5,values_only=True)}
assert ids==expected and len(ids)==57
cached=load_workbook(path,data_only=True)
assert all(cached['검토안내'][c].value==val for c,val in cache.items())

font=ImageFont.truetype('C:/Windows/Fonts/malgun.ttf',15)
def wrap(text,width):
    lines=['']
    for ch in str(text or ''):
        if ch=='\n' or (lines[-1] and font.getlength(lines[-1]+ch)>width):lines.append('')
        if ch!='\n':lines[-1]+=ch
    return lines
widths=[int(v.active.column_dimensions[c].width*7) for c in 'ABCD']
for row in v.active.iter_rows(min_row=5):
    for c,width in zip(row,widths):
        assert len(wrap(c.value,width-14))*19+10<=v.active.row_dimensions[c.row].height*4/3,c.coordinate
rows=[4]+[v.active[c].row for c in changed]
im=Image.new('RGB',(sum(widths)+20,60*len(rows)+20),'white');draw=ImageDraw.Draw(im)
for index,r in enumerate(rows):
    x=10;y=10+60*index
    for c,width in zip(v.active[r],widths):
        bg='#17324D' if r==4 else '#FFF2D5' if c.column==4 else '#EEF3F7'
        draw.rectangle((x,y,x+width,y+60),fill=bg)
        draw.multiline_text((x+7,y+9),'\n'.join(wrap(c.value,width-14)),font=font,fill='white' if r==4 else '#233D50',spacing=3);x+=width
im.save(out/'audit-preview.png')
print(json.dumps({'archive_and_xml':'OK','freeze_and_filters':'none','source_coverage':len(ids),'summary_rows':53,'source_cells':'unchanged','summary_formulas_verified':len(cache),'fixed_cells':changed,'layout':'no estimated clipping'},ensure_ascii=False))
