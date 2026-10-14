import json, pathlib, zipfile
import xml.etree.ElementTree as ET
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont

OUT = pathlib.Path(__file__).parent
DATA = json.loads((OUT / 'review_data.json').read_text(encoding='utf-8'))
DEST = OUT / '가스펠플러스_사업실적_대표경력_검토용.xlsx'
NAVY, TEAL, PALE, INK = '17324D', '087F8C', 'EAF5FA', '243746'
wb = Workbook()
guide = wb.active
guide.title = '검토안내'
guide.append(['가스펠플러스 | 사업실적·대표경력 검토대장'])
guide.merge_cells('A1:F1')
guide.append(['작성 기준: 2026-09-22 | 내부 검토용 초안 · 공개 여부 미확정'])
guide.merge_cells('A2:F2')
guide.append([])
guide.append(['검토 순서', '시트', '항목 수', '검토완료', '게시 가능', '사용 방법'])
cached = {}
for idx, (name, desc) in enumerate([
    ('회사실적_문서근거', '회사소개서·실적증명서·계약문서 기반. 실제 수행·귀속·기간을 확인하세요.'),
    ('대표경력_개인', '대표 개인 경력입니다. 이전 직장 경력을 회사 자체 실적으로 합산하지 마세요.'),
    ('사업후보_전체폴더', '폴더·파일명 기반 후보입니다. 중복·구매·제안이 포함되며 실적 건수가 아닙니다.'),
], 5):
    last = len(DATA[name]['rows']) + 4
    guide.append([idx-4, name, f"=COUNTA('{name}'!A5:A{last})", f'=COUNTIF(\'{name}\'!G5:G{last},"확인완료")', f'=COUNTIF(\'{name}\'!H5:H{last},"게시가능")', desc])
    guide.cell(idx, 2).hyperlink = f"#'{name}'!A1"
    cached[f'C{idx}'] = len(DATA[name]['rows'])
    cached[f'D{idx}'] = cached[f'E{idx}'] = 0

notes = [
    ('1. 확인사항 먼저', '확인필요사항 시트의 연도 충돌·소속·중복·게시 조건을 검토합니다.'),
    ('2. 원문과 수정값 구분', '각 목록의 A:F는 조사 내용입니다. 파란색 G:M에 검토 상태·게시 여부·확정 회사·파트·연도·문구·메모를 입력합니다.'),
    ('3. 게시할 항목 선택', '대표 검토=확인완료, 게시 여부=게시가능으로 표시한 행을 필터링합니다. 홈페이지에는 확정값과 게시용 문구를 사용합니다.'),
    ('4. 기간 해석', '파일 수정일은 참여연도로 쓰지 않았습니다. 계약기간·납품예정일·연혁일은 구별했으며 미확인 값은 빈칸 또는 확인 필요로 남겼습니다.'),
    ('5. 현재·진행중 표현', '이력서의 현재는 작성 당시 표현입니다. 지금도 수행 중이라는 뜻으로 사용하지 말고 종료일을 확인하세요.'),
    ('6. 관련 회사 구분', '관련 회사·기관은 원문 기재값입니다. 직접 발주처·최종 수요처·지원기관이 섞여 있을 수 있으니 확정 회사명에 관계를 명시하세요.'),
    ('7. 자료 근거', '각 행 우측의 원본 경로·문서 위치를 확인하세요. 폴더별_근거찾기는 후보별 참고 파일 최대 3건을 제공합니다.'),
    ('8. 조사 범위', '요청한 세 경로에서 574개 폴더 및 하위 2단계 문서 목록을 조사했습니다. 모든 파일 본문을 전수 검토한 결과는 아닙니다. 휴지통·백업·설치자료 등은 제외했습니다.'),
    ('9. 확정 실적과 후보', '문서근거 시트에도 계약 초안·납품 예정 항목이 포함됩니다. 근거 수준과 확인사항을 보고 완료 여부를 확정하세요.'),
    ('10. 동일 사업 중복', '회사·개인 경력·폴더 후보 간 중복이 있습니다. 서로 더해 총 수행 건수로 사용하지 마세요.'),
    ('11. 국방·항공 및 공개 조건', '대표 검토 후 사업명·고객명·담당 파트의 공개 범위를 확정하세요. 오스템 계약문서의 공개 관련 조항은 확인필요사항 Q12에 표시했습니다.'),
    ('12. 개인정보 제외', '연락처·주소·생년월일·계좌·계약금액은 이 검토대장에 옮기지 않았습니다.'),
    ('13. 분류 기준', '분야는 검토 편의를 위한 분류입니다. 폴더 기반 회사명·참여 파트는 단서이며 실제 수행을 확정한 값이 아닙니다.'),
    ('14. 기타 시트', '대표_재직이력: 원문 재직기간·직위와 차이. 확인필요사항: 대표 확인 요청. 폴더별_근거찾기: 원본 폴더·참고 문서 목록.'),
]
guide.append([])
for label, text in notes:
    r = guide.max_row + 1
    guide.cell(r, 1, label)
    guide.cell(r, 2, text)
    guide.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)

widths_main = [11, 23, 43, 28, 49, 34, 16, 16, 29, 47, 25, 60, 45, 41, 29, 62, 13, 13, 84, 37]
order = [0,2,1,3,4,5,13,14,15,16,17,18,19,8,9,12,6,7,10,11]

def validation(ws, col, choices, last):
    dv = DataValidation(type='list', formula1='"'+','.join(choices)+'"', allow_blank=True)
    dv.errorTitle, dv.error = '목록에서 선택', '드롭다운에 있는 값을 선택해 주세요.'
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(f'{col}5:{col}{last}')

for index, (name, data) in enumerate(DATA.items(), 1):
    ws = wb.create_sheet(name)
    main = len(data['headers']) == 20
    headers = [data['headers'][i] for i in order] if main else data['headers']
    rows = [[row[i] for i in order] for row in data['rows']] if main else data['rows']
    ws.append([name.replace('_', ' | ')])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(len(headers), 8))
    ws.append(['파란색 칸에 대표 검토·확정값을 입력하세요. 원본 경로와 확인사항은 오른쪽 열에 있습니다.' if main else '내부 검토용 | 원문 근거와 확인할 내용을 함께 보관합니다.'])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=min(len(headers), 8))
    ws.append([])
    ws.append(headers)
    for row in rows:
        ws.append(row)
    widths = widths_main if main else {
        '확인필요사항':[11,26,30,65,58,90,18,55],
        '대표_재직이력':[30,26,25,20,26,67,95],
        '폴더별_근거찾기':[12,50,18,43,95,110,54,17],
    }[name]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    table = Table(displayName=f'ReviewTable{index}', ref=f'A4:{get_column_letter(len(headers))}{ws.max_row}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = 'D5' if main else 'C5'
    ws.sheet_view.zoomScale = 75 if main else 85
    ws.sheet_properties.tabColor = TEAL if main else '758A99'
    ws.print_title_rows = '1:4'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_area = f'A1:{"M" if main else get_column_letter(ws.max_column)}{ws.max_row}'
    if main:
        validation(ws, 'G', ['미검토','확인완료','수정필요','중복','제외'], ws.max_row)
        validation(ws, 'H', ['미정','게시가능','익명게시','비공개','보류'], ws.max_row)
        for col, good, bad in [('G','확인완료','수정필요'),('H','게시가능','비공개')]:
            for val, fill in [(good,'DDF0DF'),(bad,'FFE1D7')]:
                ws.conditional_formatting.add(f'{col}5:{col}{ws.max_row}',CellIsRule(operator='equal',formula=['"'+val+'"'],fill=PatternFill('solid',fgColor=fill)))
        for row in ws.iter_rows(min_row=5, min_col=7, max_col=13):
            for cell in row:
                cell.fill = PatternFill('solid', fgColor=PALE)
                cell.font = Font(name='맑은 고딕',size=11,color='005C99')
    elif name == '확인필요사항':
        validation(ws, 'G', ['미검토','확인완료','추가자료필요','보류'], ws.max_row)
        for row in ws.iter_rows(min_row=5,min_col=7,max_col=8):
            for cell in row: cell.fill = PatternFill('solid',fgColor=PALE)

for i,width in enumerate([27,34,14,14,14,70],1):
    guide.column_dimensions[get_column_letter(i)].width=width
guide.freeze_panes='C5'
guide.sheet_properties.tabColor=NAVY

font = ImageFont.truetype('C:/Windows/Fonts/malgun.ttf', 15)
bold = ImageFont.truetype('C:/Windows/Fonts/malgunbd.ttf', 15)

def wrap(text, pixel_width, use_font=font):
    lines=[]
    for para in str('' if text is None else text).split('\n'):
        line=''
        for char in para:
            if line and use_font.getlength(line+char)>pixel_width:
                lines.append(line);line=''
            line+=char
        lines.append(line)
    return lines

for ws in wb:
    ws.sheet_view.showGridLines=False
    for row in ws:
        for cell in row:
            if cell.__class__.__name__ == 'MergedCell': continue
            if not cell.font.color or cell.font.color.type != 'rgb':
                cell.font=Font(name='맑은 고딕',size=11,color=INK)
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            if cell.data_type == 'e': raise ValueError(f'Error cell {ws.title}!{cell.coordinate}')
            if cell.row==4:
                cell.fill=PatternFill('solid',fgColor=TEAL if 7<=cell.column<=13 and ws.max_column==20 else NAVY)
                cell.font=Font(name='맑은 고딕',bold=True,size=11,color='FFFFFF')
            if cell.row==1:
                cell.font=Font(name='맑은 고딕',bold=True,size=20,color='FFFFFF')
                cell.fill=PatternFill('solid',fgColor=NAVY)
            if cell.row==2:
                cell.font=Font(name='맑은 고딕',size=11,color='526A7D')
    ws.row_dimensions[1].height=38
    ws.row_dimensions[2].height=32
    ws.row_dimensions[3].height=10
    ws.row_dimensions[4].height=36
    for r in range(5,ws.max_row+1):
        longest=1
        for cell in ws[r]:
            if cell.__class__.__name__ == 'MergedCell': continue
            width=ws.column_dimensions[cell.column_letter].width*7
            if ws==guide and r>=9 and cell.column==2:
                width=sum(ws.column_dimensions[get_column_letter(c)].width*7 for c in range(2,7))
            longest=max(longest,len(wrap(cell.value,width-14)))
        ws.row_dimensions[r].height=max(42,min(360,longest*13+12))
    ws.auto_filter.ref=None  # Excel table owns the filter range.

wb.save(DEST)
# Cache the guide's six COUNTIFs and three COUNTA formulas for preview readers.
ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
with zipfile.ZipFile(DEST) as z:
    parts={item.filename:z.read(item.filename) for item in z.infolist()}
root=ET.fromstring(parts['xl/worksheets/sheet1.xml'])
for cell in root.findall('.//s:c',ns):
    if cell.attrib['r'] in cached:
        cell.find('s:v',ns).text=str(cached[cell.attrib['r']])
parts['xl/worksheets/sheet1.xml']=ET.tostring(root,encoding='utf-8')
with zipfile.ZipFile(DEST,'w',zipfile.ZIP_DEFLATED) as z:
    for name,content in parts.items(): z.writestr(name,content)

check=load_workbook(DEST)
values=load_workbook(DEST,data_only=True)
assert len(check.worksheets)==7
for name,data in DATA.items():
    ws=check[name]
    assert ws.max_row==len(data['rows'])+4
    expected=[[r[i] for i in order] for r in data['rows']] if len(data['headers'])==20 else data['rows']
    for actual,row in zip(ws.iter_rows(min_row=5,values_only=True),expected):
        assert list(actual)==[None if x=='' else x for x in row],(name,actual[0])
    if len(data['headers'])==20:
        assert len(ws.data_validations.dataValidation)==2
        assert ws['G5'].value=='미검토' and ws['H5'].value=='미정'
        assert ws['Q5'].value is None or isinstance(ws['Q5'].value,int)
for addr,val in cached.items(): assert values['검토안내'][addr].value==val
assert not any(c.data_type=='e' for ws in check for row in ws for c in row)

# Render representative ranges from the saved workbook's cell styles and dimensions.
previews=[]
for idx,ws in enumerate(values):
    cols=min(ws.max_column,6)
    rows=list(range(1,min(ws.max_row,8)+1))
    widths=[int(ws.column_dimensions[get_column_letter(c)].width*7) for c in range(1,cols+1)]
    heights=[int(ws.row_dimensions[r].height*1.333) for r in rows]
    img=Image.new('RGB',(sum(widths)+32,sum(heights)+55),'white')
    draw=ImageDraw.Draw(img)
    draw.text((16,8),ws.title,font=bold,fill='#17324D')
    y=38
    for r,h in zip(rows,heights):
        x=16
        for c,w in enumerate(widths,1):
            cell=ws.cell(r,c)
            if cell.__class__.__name__=='MergedCell':
                x+=w;continue
            span=next((m for m in ws.merged_cells.ranges if m.min_row==r and m.min_col==c),None)
            cw=sum(widths[c-1:min(span.max_col,cols)]) if span else w
            bg='#'+cell.fill.fgColor.rgb if cell.fill.fgColor.type=='rgb' and cell.fill.patternType=='solid' else ('#F1F5F8' if r>=5 and r%2==1 else 'white')
            draw.rectangle((x,y,x+cw,y+h),fill=bg)
            fg='#'+cell.font.color.rgb[-6:] if cell.font.color and cell.font.color.type=='rgb' else '#243746'
            f=bold if r in [1,4] else font
            text='\n'.join(wrap(cell.value,cw-14,f))
            draw.multiline_text((x+7,y+6),text,font=f,fill=fg,spacing=2)
            x+=w
        y+=h
    p=OUT/f'preview-{idx}.png';img.save(p);previews.append(str(p))
print(json.dumps({'file':str(DEST.resolve()),'sheets':{ws.title:ws.max_row-4 for ws in check if ws.title!='검토안내'},'validation':'All cell values, counts, dropdowns and formula caches verified','previews':previews},ensure_ascii=False,indent=2))
