import pathlib
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont

OUT=pathlib.Path(__file__).parent
path=OUT/'가스펠플러스_사업실적_대표경력_검토용.xlsx'
wb=load_workbook(path)
sources={r[0]:r for name in ['회사실적_문서근거','대표경력_개인'] for r in wb[name].iter_rows(min_row=5,values_only=True)}
# Each group is an editorial case bundle, not a claim that the contracts are identical.
groups=[
('먼저 검토','C021 C022 P001 P002','의료기기 PCB·제품 기능검사장비 개발','안마기 PCB와 제품 부품의 기능검사를 위한 장비를 설계·제작했습니다. 전기회로 설계부터 검사 프로그램과 GUI 개발, 장비 제작 및 사후 지원까지 참여했습니다.','PCB 검사장비 제작, 기능검사 지그, 의료기기 검사지그','회사 사례','V9·V11의 개별 납품 시점과 검사장비 종수, 고객명 공개 범위 확인.'),
('먼저 검토','C027','다중 생체신호 측정용 웨어러블 센서 보드 개발','피로도와 호흡 분석에 활용하는 웨어러블 센서 보드를 개발했습니다. 다중 생체신호 측정 보드와 시험 펌웨어, 회로도 및 설명문서를 제공했습니다.','웨어러블 센서 개발, 생체신호 측정 보드, 센서 PCB 개발','회사 사례','실적증명서에 이행 100% 기재. 고객명 및 소개문 공개 범위 확인.'),
('먼저 검토','C006 C009','LED 전원 드라이버·가시광통신 제어 개발','소형 LED 전등과 가로등급 LED 전원 드라이버를 개발했습니다. LED 구동·전원 제어와 가시광통신 신호처리를 위한 제어보드 개발에 참여했습니다.','LED 드라이버 개발, LED 전원회로 설계, 가시광통신 보드','회사 사례','소형 LED와 가로등 드라이버는 별도 사례. 계약기간과 소개서 연혁 시점을 구분해 확정.'),
('먼저 검토','C015 C020','드론 전원·전력분배·통신·센싱 보드 개발','드론용 파워보드와 전력분배보드를 개발했습니다. 후속 사례로 파워·통신·센싱 보드 개발에 참여한 경험이 있습니다.','드론 전원보드 개발, PDB 개발, 드론 통신보드 제작','회사 사례','2020년·2024년 사업을 구분하고 고객 및 실제 개발 범위 확인.'),
('먼저 검토','C014','다채널 계측·디스플레이 보드 개발','32채널 측정 및 디스플레이 보드를 개발했습니다. 여러 채널의 신호를 다루는 계측·표시 보드 개발 사례입니다.','다채널 계측보드 개발, 아날로그 입력보드, 디스플레이 보드 제작','회사 사례','측정 신호 종류·납품처·담당 범위 확인. 확인되지 않은 정밀도는 기재하지 않음.'),
('먼저 검토','C005','Zigbee 무선 센서 노드 보드 개발','Zigbee Node 보드를 개발했습니다. 무선통신을 적용한 임베디드 보드 개발 사례입니다.','Zigbee 보드 개발, 무선통신 PCB 설계, IoT 노드 개발','회사 사례','펌웨어 포함 여부·통신 구성·실제 활용처 확인.'),
('먼저 검토','C018 P008','이동형 VR 경찰훈련장비 기구·전장 제작','이동형 VR 경찰훈련 시뮬레이터의 기구·전장 설계와 제작에 참여했습니다. 훈련장비 하드웨어 제작 사례입니다.','VR 훈련장비 제작, 시뮬레이터 전장 설계, 훈련장비 하드웨어','회사 사례','회사소개서와 대표 이력서의 관련 사례를 묶음. 소프트웨어 전체 개발로 표현하지 않음.'),
('먼저 검토','C019 P007','낙하산 시뮬레이터 전장 제작·설치','낙하산 시뮬레이터의 시스템 전장을 제작하고 설치했습니다. 훈련 시스템에 필요한 전장 제작·설치 참여 사례입니다.','시뮬레이터 전장 제작, 훈련장비 배선, 전장 설치','회사 사례','군 관련 사업명·고객명 공개 범위와 실제 담당 부분 확인.'),
('완료 확인 후','C026','광자극·망막 전위차 측정 시스템 개발','고속·고휘도 광자극과 망막 전위차 측정을 위한 시스템 개발 과제에 참여했습니다. LCD 백라이트와 제어보드, ERG 측정보드, 펌웨어 및 시험이 계약 범위에 포함됩니다.','ERG 측정장비 개발, 광자극기 개발, 생체신호 계측 시스템','회사 계약 근거','계약기간은 확인됨. 실제 검수·완료 여부 확인 후 완료실적으로 표현.'),
('완료·공개 확인 후','C028','의료장비 PCB 4종 기능검사장비 제작','의료장비용 PCB 4종의 검사장비 제작·납품을 대상으로 한 사례입니다. 보드별 기능을 확인하는 전용 검사장비 구축 범위를 검토했습니다.','의료장비 PCB 검사, PCB 기능검사장비 제작, 전용 검사지그','계약문서 후보','서명된 계약서·실제 수행·납품 확인 전 확정 실적으로 게시하지 않음. 원문 공개 조항 확인.'),
('귀속·기간 확인 후','C012 P011','반도체 웨이퍼 공정온도 모니터링 시스템','반도체 공정의 웨이퍼 온도 모니터링을 위한 무선통신 시스템 개발에 참여했습니다. 임베디드 보드 설계, 펌웨어와 Android 앱 개발이 소개자료에 기재돼 있습니다.','웨이퍼 온도 측정, 공정 모니터링 개발, 무선 온도 계측 시스템','회사·개인 귀속 확인','회사소개서와 모비다임 재직 중 개인 경력에 중복 기재. 회사 실적 또는 대표 경력으로 확정 필요.'),
('귀속·기간 확인 후','C016 P014','복강경용 마그네틱 마킹·위치 측정 장비 개발','마그네틱 센서를 이용한 복강경 마킹·거리 측정 장비 개발에 참여했습니다. 임베디드 보드 설계와 펌웨어, OS 포팅 및 Android 앱 개발이 소개자료에 기재돼 있습니다.','복강경 장비 개발, 마그네틱 센서 개발, 의료기기 임베디드 개발','회사·개인 귀속 확인','2020년과 2021년 표기가 다름. 회사 귀속·개발 단계·기간 확정 필요.'),
('귀속·기간 확인 후','C008 C024 P016','자동차 휠 속도센서·검증 지그 개발','고해상도 휠 속도센서와 센서 검증 지그 개발 경험을 보유하고 있습니다. 임베디드 보드와 펌웨어, 검증용 제어 앱·GUI 개발이 소개자료에 기재돼 있습니다.','휠 속도센서 개발, 자동차 센서 검증 지그, 홀센서 측정장비','회사·개인 귀속 확인','WSS·검증 지그·개인 홀센서 프로젝트는 관련 사례 묶음. 동일 계약으로 단정하지 않음. GM/중기청 관계 확인.'),
('중복·기간 확인 후','C023 C025 P015','무선 체온 측정·BLE 센서 장비 개발','체온 측정과 BLE 데이터 전송을 위한 무선 센서 장비 개발 경험을 보유하고 있습니다. 관련 소개자료에는 보드 설계·펌웨어·Android 앱 개발이 기재돼 있습니다.','BLE 체온계 개발, 무선 온도센서 개발, 의료용 센서 보드','회사·개인 귀속 확인','미래테크원 온도계·신생아 온도센서·소개서 체온계는 별도일 수 있음. 기간·고객·개발범위 연결 확인.'),
('추가 정보 확인','C003','커피 로스터 전장 개발','커피 로스터의 전장 개발에 참여했습니다. 제품의 전기·전자 구성 개발 경험을 소개하는 사례입니다.','커피 로스터 전장 개발, 가전 전자제어 개발, 임베디드 제품 개발','회사 사례','고객명·제어 기능·회로 및 펌웨어 분담 확인.'),
('추가 정보 확인','C004','신발건조기 개발','신발건조기 개발에 참여했습니다. 제품 개발 과정에서 맡은 세부 범위를 확인해 소개할 사례입니다.','신발건조기 개발, 생활가전 개발','회사 사례','세부 참여 파트가 확인되지 않아 회로·기구·펌웨어 전체 개발로 쓰지 않음.'),
('추가 정보 확인','C011','태양광 무선 수로 개폐장치 개발','태양광을 활용한 무선 수로 개폐장치 개발에 참여했습니다. 수로 관리 장치에 적용한 개발 사례입니다.','무선 수로 개폐장치, 태양광 제어장치 개발, 농업용 제어장치','회사 사례','실제 담당 회로·제어·통신 범위와 발주처 확인.'),
('추가 정보 확인','C017','오토바이 무선충전 거치대 개발·양산','오토바이용 무선충전 거치대를 개발하고 양산에 참여했습니다. 제품 개발에서 생산까지 이어진 사례입니다.','무선충전 거치대 개발, 오토바이 무선충전기, 전자제품 양산 개발','회사 사례','고객명·개발 분담·양산 범위 확인.'),
('공개·범위 확인','C007 C013','헬기 훈련장비 패널 성능개선','LYNX 계열 훈련장비의 성능개선 사업과 패널류 성능개선에 참여했습니다. 훈련장비 유지·개선을 위한 하드웨어 참여 사례입니다.','시뮬레이터 패널 개발, 훈련장비 성능개선, 항공 계기패널','회사 사례','2017년 CFT와 2019년 패널류 개선을 구분. 원문 LYXY 표기와 공개 가능한 사업명 확인.'),
('공개·범위 확인','C010','잠수함 훈련장비 패널 개발','잠수함 시뮬레이터용 패널 7종 개발에 참여했습니다. 훈련장비의 패널 개발 사례입니다.','훈련장비 패널 개발, 시뮬레이터 계기패널, 전자 패널 제작','회사 사례','2018년 회사 사례. 대표의 이전 잠수함 경력과 별도 구분 및 공개 범위 확인.'),
('귀속·기간 확인 후','C001 P020','위성 관련 GRA 시뮬레이터·데이터 장치 개발','GRA 시뮬레이터와 위성 관련 데이터 저장·전송장치 개발 경험이 자료에 기재돼 있습니다. 세부 사업과 담당 업무를 확인해 항공우주 개발 사례로 소개할 후보입니다.','위성 시험장비 개발, 데이터 저장 전송 보드, GRA 시뮬레이터','회사·개인 귀속 확인','회사 GRA 시뮬레이터와 개인 GRA 장치가 같은 사업인지 미확정. 경력표 날짜 역전 오류 확인.'),
('귀속·기간 확인 후','C002 P017','위성 시험 인터페이스 보드 개발','다목적실용위성 6호 관련 Baseband Test Interface Unit 보드 개발에 참여한 경험이 있습니다. 위성 시험용 인터페이스 보드 개발 사례입니다.','위성 시험 인터페이스 보드, TIU 보드 개발, 항공우주 PCB 개발','회사·개인 귀속 확인','회사소개서 2016.12 / 이력서 2019. 개발 차수·수행회사·기간 확인.'),
('대표 경력으로 검토','P003','6축 모션 시스템 서보랙 제작 참여','대표 진형보는 6축 모션 시스템의 서보랙 제작에 PL로 참여한 경력이 있습니다.','서보랙 제작, 모션 시스템 전장, 6축 모션 전기제어','대표 개인 경력','에레모스 발주·LIG 사업명 구분. 회사 계약 근거 확인 시 회사 사례로 전환 가능.'),
('대표 경력으로 검토','P004','항공용 멀티기능모니터 개발 참여','대표 진형보는 회전익·고정익 항공용 멀티기능모니터 CDNU·MFCU 개발에 PL로 참여한 경력이 있습니다.','항공 디스플레이 개발, 멀티기능모니터 개발, 임베디드 모니터','대표 개인 경력','2023~2024년 소속·회사 계약 주체·세부 개발 범위 확인.'),
('대표 경력으로 검토','P005','원자력장비 전장·케이블 제작 참여','대표 진형보는 원자력장비 전장 작업과 케이블 제작에 PL로 참여한 경력이 있습니다.','원자력장비 전장, 산업용 케이블 제작, 장비 배선 제작','대표 개인 경력','회사 계약 주체와 수행 범위 확인 시 회사 사례로 전환 가능.'),
('대표 경력으로 검토','P006','방호복 착탈의 훈련장비 개발·설치 참여','대표 진형보는 감염병 방호복 착탈의 훈련장비 개발·설치에 PL로 참여한 경력이 있습니다.','의료 훈련장비 제작, 방호복 훈련장비, 교육장비 개발','대표 개인 경력','회사 계약 주체·전장 또는 전체 시스템 분담 확인.'),
('대표 경력으로 검토','P009','차량 공조 튜닝·MCU 업데이트 장치 개발 참여','대표 진형보는 차량 공조장치 로직 튜닝과 MCU ROM 업데이트 장치 개발에 PL로 참여한 경력이 있습니다.','차량 공조 시험장비, MCU 업데이트 장치, 자동차 제어 검증','대표 이전 직장 경력','모비다임 재직 중 경력. 회사 직접 실적으로 표현하지 않음.'),
('대표 경력으로 검토','P010','차량 베어링 진동 감지 시스템 개발 참여','대표 진형보는 차량용 베어링 진동 자동 감지 시스템 개발에 PL로 참여한 경력이 있습니다.','베어링 진동 감지, 자동차 진동센서, 센서 모니터링 시스템','대표 이전 직장 경력','최종 수요처와 계약 상대 구분. 모비다임 재직 중 경력.'),
('대표 경력으로 검토','P012','의료용 pH 측정 밴드 개발 참여','대표 진형보는 의료용 환부 pH 측정 밴드 장비 개발에 PL로 참여한 경력이 있습니다.','pH 센서 보드 개발, 의료용 센서 밴드, 웨어러블 측정 장비','대표 이전 직장 경력','모비다임 재직 중 경력. 본인 담당 개발 범위를 추가 확인.'),
('대표 경력으로 검토','P013','압력센싱 보드 개발 참여','대표 진형보는 압력센싱 보드 개발에 PL로 참여한 경력이 있습니다.','압력센서 보드 개발, 압력 측정 PCB, 자동차 센싱 보드','대표 이전 직장 경력','모비다임 재직 중 경력. 고객 공개 범위 확인.'),
('대표 경력으로 검토','P019','곡률 측정용 로드셀 장비 개발 참여','대표 진형보는 커브모니터 제작 공정의 곡률 측정을 위한 로드셀 장비 개발에 참여한 경력이 있습니다.','로드셀 계측장비 개발, 공정 측정장비, 곡률 측정장비','대표 이전 직장 경력','IMNS 재직 중 경력. 담당 회로·소프트웨어 범위 확인.'),
('대표 경력으로 검토','P021','항공기 간헐결함 탐지장비 개발 참여','대표 진형보는 항공기 간헐결함 탐지장비의 전기전자 개발에 참여한 경력이 있습니다.','간헐결함 탐지장비, 항공 정비 시험장비, 전기전자 시험장비','대표 이전 직장 경력','2014~2016 / 2014~2015 기간 차이와 공개 명칭 확인.'),
('대표 경력으로 검토','P022 P024 P025 P026','항공 훈련 시뮬레이터 계기패널·전기전자 개발 참여','대표 진형보는 헬기·수송기 시뮬레이터의 계기패널과 비행훈련장비 전기전자 장치 개발에 참여한 경력이 있습니다.','비행훈련장비 개발, 항공 계기패널, 시뮬레이터 전기전자','대표 이전 직장 경력','UH60·LYNX·C130·KT100 개별 사례 묶음. 사업별 기간 충돌과 공개 범위 확인.'),
('대표 경력으로 검토','P023','잠수함 조종훈련장비 계기패널 개발 참여','대표 진형보는 잠수함 시뮬레이터 계기패널 개발에 참여한 경력이 있습니다.','조종훈련장비 계기패널, 시뮬레이터 패널 설계','대표 이전 직장 경력','이력서 2011 / 경력표 2012. 회사의 2018년 패널 개발 사례와 구분.'),
('대표 경력으로 검토','P018','ACMI POD 전기전자 개발 참여','대표 진형보는 ACMI POD의 전기전자 개발에 참여한 경력이 있습니다.','항공 전자장치 개발, ACMI POD 전기전자','대표 이전 직장 경력','이력서의 제품 분류와 고객 LIG/ADD·LIG Nex 표기 및 공개 범위 확인.'),
('대표 경력으로 검토','P027 P028','모션 시스템·로봇 구동장치 전장 설계·제작 참여','대표 진형보는 6축 모션 시스템과 구난로봇 Actuator 개발사업에서 전장 설계·제작에 참여한 경력이 있습니다.','모션 시스템 전장 설계, 로봇 전장 제작, 액추에이터 제어장비','대표 이전 직장 경력','삼정 6축 모션과 구난로봇은 별도 사업. IMNS 재직 중 경력.'),
('진행·체결 확인 후','C029','광제어용 생체신호 측정모듈 개발','광제어를 위한 생체신호 측정모듈 개발·현장설치를 대상으로 한 과제입니다. 계약 체결과 실제 수행 상태를 확인한 뒤 진행 과제 또는 완료 사례로 소개할 후보입니다.','생체신호 측정모듈 개발, 광제어 계측보드','계약 초안 후보','초안만 확인됐으며 납품 예정일은 2027.02.26. 완료실적으로 게시하지 않음.'),
]
assert set(' '.join(g[1] for g in groups).split())==set(sources)
if '홈페이지_한장정리' in wb: del wb['홈페이지_한장정리']
ws=wb.create_sheet('홈페이지_한장정리',0)
for n,text in enumerate([
 '가스펠플러스 | 홈페이지에 보여줄 개발·제작 사례',
 '이 시트 하나에서 검토하세요. 제목·소개문을 읽고 G열에서 결정, H열에 수정 내용을 적으면 됩니다. 모두 게시 전 검토용 초안입니다.',
 '관련 사례를 주제별로 묶었습니다. 한 행이 한 계약을 뜻하지 않습니다. 회사·대표 이전 직장 경력은 F열에서 구분합니다.',
 '검색어는 업무 내용을 풀어 쓴 문안 후보이며 검색량 분석 결과가 아닙니다. 고객명은 소개문에서 제외했고 원문·기간은 E열과 메모에 남겼습니다.',
 '기존 조사자료와 574개 폴더 후보는 같은 파일의 숨김 시트에 보관했습니다. 필요한 경우 Excel 시트 탭 → 숨기기 취소로 확인하세요.'
],1):
    ws.cell(n,1,text);ws.merge_cells(start_row=n,start_column=1,end_row=n,end_column=8)
headers=['검토 순서','홈페이지 제목 초안','홈페이지 소개문 초안','검색어·문의 키워드 후보','관련 사업·참여연도·회사·파트 (원문)','실적 구분 / 게시 전 확인','대표 결정','대표 수정 메모','원본 항목 ID','원본 경로']
ws.append(headers)
for priority,ids,title,copy,keywords,kind,warning in groups:
    refs=[sources[i] for i in ids.split()]
    facts='\n\n'.join(f'{r[0]} {r[2]}\n기간: {r[5]}\n회사·기관: {r[3]}\n파트: {r[4]}' for r in refs)
    paths='\n\n'.join(dict.fromkeys(str(r[18])+' | '+str(r[19]) for r in refs))
    ws.append([priority,title,copy,keywords,facts,kind+'\n'+warning,'미검토',None,ids,paths])
    row=ws.max_row
    ws.cell(row,2).comment=Comment('근거 항목: '+ids+'\n'+paths,'자료 정리')
    ws.cell(row,6).comment=Comment('\n\n'.join(r[0]+': '+str(r[15] or '추가 확인사항 없음') for r in refs),'자료 정리')
    for r in refs:
        if any(r[i] not in (None,'미검토','미정') for i in range(6,13)):
            ws.cell(row,8).value=(ws.cell(row,8).value or '')+'\n기존 검토 '+r[0]+': '+str(r[6:13])
widths=[22,43,78,44,73,74,20,60,28,100]
for i,w in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=w
ws.column_dimensions.group('I','J',hidden=True)
for row in ws:
    for cell in row:
        if cell.__class__.__name__=='MergedCell':continue
        cell.font=Font(name='맑은 고딕',size=11,color='233D50')
        cell.alignment=Alignment(vertical='top',wrap_text=True)
        if cell.row==1:
            cell.font=Font(name='맑은 고딕',size=20,bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17324D')
        elif cell.row==6:
            cell.font=Font(name='맑은 고딕',size=11,bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='087F8C' if cell.column in (7,8) else '17324D')
        elif cell.row>=7 and cell.column in (7,8):
            cell.fill=PatternFill('solid',fgColor='E8F4FB');cell.font=Font(name='맑은 고딕',size=11,color='005C99')
ws.row_dimensions[1].height=42
for r in range(2,6):ws.row_dimensions[r].height=25
ws.row_dimensions[6].height=40
# Long original evidence is retained in a cell comment; visible facts stay readable.
for r in range(7,ws.max_row+1):
    full=ws.cell(r,5).value
    ws.cell(r,5).comment=Comment(full,'원문 대조')
    refs=[sources[i] for i in ws.cell(r,9).value.split()]
    compact='\n'.join(f'{x[0]} | {x[5]} | {x[3]}' for x in refs)
    roles=list(dict.fromkeys(x[4] for x in refs))
    ws.cell(r,5).value=compact+'\n담당: '+' / '.join(roles)
    ws.row_dimensions[r].height=max(120,36+len(refs)*31+len(' / '.join(roles))//35*16)
table=Table(displayName='HomepageReview',ref=f'A6:H{ws.max_row}')
table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
ws.add_table(table)
dv=DataValidation(type='list',formula1='"미검토,게시후보,수정필요,보류,제외"')
dv.showErrorMessage=True;dv.error='목록에서 선택해 주세요.';ws.add_data_validation(dv);dv.add(f'G7:G{ws.max_row}')
ws.conditional_formatting.add(f'G7:G{ws.max_row}',CellIsRule(operator='equal',formula=['"게시후보"'],fill=PatternFill('solid',fgColor='DDF0DF')))
ws.freeze_panes='C7';ws.sheet_view.zoomScale=65;ws.sheet_view.showGridLines=False
ws.print_title_rows='1:6';ws.print_area=f'A1:H{ws.max_row}'
ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
ws.sheet_properties.pageSetUpPr.fitToPage=True
for sheet in list(wb)[1:]:sheet.sheet_state='hidden'
wb.active=0;wb.save(path)
check=load_workbook(path)
assert [s.title for s in check if s.sheet_state=='visible']==['홈페이지_한장정리']
assert check.active.max_row==len(groups)+6
for name in ['회사실적_문서근거','대표경력_개인']:
    for row in check[name].iter_rows(min_row=5,values_only=True):assert row==sources[row[0]]
assert all(check.active.cell(r,7).value=='미검토' for r in range(7,check.active.max_row+1))

# Visual QA of saved cell content, wrapping and review fields.
font=ImageFont.truetype('C:/Windows/Fonts/malgun.ttf',16)
def wrap(text,width):
    lines=[]
    for p in str(text or '').split('\n'):
        line=''
        for ch in p:
            if line and font.getlength(line+ch)>width:lines.append(line);line=''
            line+=ch
        lines.append(line)
    return lines
for name,cols,rows in [('main',[1,2,3,4],[6,7,8,9]),('review',[5,6,7,8],[6,7,8,9]),('career',[2,3,6],[6,39,40,41])]:
    widths_px=[int(widths[c-1]*7) for c in cols]
    hs=[max(60,max(len(wrap(check.active.cell(r,c).value,w-20)) for c,w in zip(cols,widths_px))*21+22) for r in rows]
    img=Image.new('RGB',(sum(widths_px)+20,sum(hs)+20),'white');draw=ImageDraw.Draw(img);y=10
    for r,h in zip(rows,hs):
        x=10
        for c,w in zip(cols,widths_px):
            bg='#17324D' if r==6 else '#E8F4FB' if c in (7,8) else '#F1F5F8' if r%2 else 'white'
            draw.rectangle((x,y,x+w,y+h),fill=bg)
            draw.multiline_text((x+10,y+10),'\n'.join(wrap(check.active.cell(r,c).value,w-20)),font=font,fill='white' if r==6 else '#233D50',spacing=3)
            x+=w
        y+=h
    img.save(OUT/f'consolidated-{name}.png')
print(f'Updated existing workbook: {len(groups)} editorial case groups; all {len(sources)} source rows preserved; one visible sheet.')
