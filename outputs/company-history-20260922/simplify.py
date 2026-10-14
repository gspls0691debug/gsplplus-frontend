import pathlib, re
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image, ImageDraw, ImageFont

out=pathlib.Path(__file__).parent
path=out/'가스펠플러스_사업실적_대표경력_검토용.xlsx'
wb=load_workbook(path)
source={r[0]:r for n in ['회사실적_문서근거','대표경력_개인'] for r in wb[n].iter_rows(min_row=5,values_only=True)}
periods={
 'C001':'2015~2016','C002':'2016 / 2019 (확인 필요)',
 'C006':'2016 계약 / 2017 연혁','C009':'2016 계약 / 2018 연혁',
 'C012':'2019 (기간·귀속 확인)','C016':'2020 / 2021 (확인 필요)',
 'C023':'확인 필요','C024':'확인 필요','C025':'2016 / 2021 (확인 필요)',
 'C026':'2025~2026 (계약기간)','C028':'2025 (납품예정 기준)',
 'C029':'시작 미확인 / 2027 납품예정',
 'P001':'2025~종료 미확인','P020':'2015 (세부기간 확인)',
 'P021':'2014~2016 (확인 필요)','P022':'2011 / 2012 (확인 필요)',
 'P023':'2011 / 2012 (확인 필요)','P024':'2008~2010 (시작연도 확인)',
 'P025':'2007~2008 (세부기간 확인)','P026':'2015~종료 미확인',
}
role_short={
 'C006':'LED 전등·전원제어 개발','C021':'회로·검사 프로그램 설계, 장비 제작',
 'C022':'회로·검사 프로그램 설계, 장비 제작',
 'C026':'광자극·ERG 측정보드, 펌웨어 개발 (계약 범위)',
 'C027':'생체신호 보드·시험 펌웨어 개발',
 'C028':'PCB 4종 검사장비 제작·납품 (수행 확인)',
 'C029':'측정모듈 개발·설치 (계약 초안)',
 'P018':'전기전자 개발','P020':'전기전자 개발','P021':'전기전자 개발',
 'P022':'계기패널 개발','P023':'계기패널 개발','P024':'계기패널 개발','P025':'계기패널 개발',
}
duplicate={'P001':'C022','P002':'C021','P007':'C019','P008':'C018'}
if '참여사업_요약' in wb: del wb['참여사업_요약']
ws=wb.create_sheet('참여사업_요약',0)
for r,text in enumerate(['참여 사업 요약','사업명 · 참여 파트 · 참여연도 | PL = 프로젝트 리더','확인 필요한 기간은 표시했습니다. 원문·세부 근거는 셀 메모에 있습니다.'],1):
    ws.cell(r,1,text);ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=4)
ws.append(['구분','사업명','참여 파트','참여연도'])
rows=[]
for ident,r in source.items():
    if ident in duplicate: continue
    kind='회사 자료' if ident.startswith('C') else '대표 개인 경력'
    if ident in ['C008','C012','C016','C024']:kind='회사 귀속 확인'
    if ident in ['C028','C029']:kind='수행 확인 필요'
    years=re.findall(r'(?:19|20)\d{2}',r[5])
    period=periods.get(ident, '~'.join(dict.fromkeys(years)) if years else '확인 필요')
    role=r[10] or role_short.get(ident,r[4])
    if role=='PL':role='프로젝트 리더(PL)'
    if r[11]:period=r[11]
    rows.append((0 if ident.startswith('C') else 1,int(years[0]) if years else 9999,ident,[kind,r[2],role,period]))
for _,_,ident,row in sorted(rows):
    ws.append(row)
    r=source[ident]
    refs=[r]+[source[k] for k,v in duplicate.items() if v==ident]
    note='\n\n'.join(f'{x[0]} | {x[3]}\n원문 기간: {x[5]}\n귀속: {x[13]}\n확인사항: {x[15] or "없음"}\n{x[18]}\n{x[19]}' for x in refs)
    ws.cell(ws.max_row,2).comment=Comment(note,'원문 근거')
    ws.cell(ws.max_row,4).comment=Comment(note,'기간 확인')
widths=[22,62,65,44]
for col,width in zip('ABCD',widths):ws.column_dimensions[col].width=width
for row in ws:
    for c in row:
        if c.__class__.__name__=='MergedCell':continue
        c.font=Font(name='맑은 고딕',size=11,color='233D50')
        c.alignment=Alignment(vertical='center',wrap_text=True)
        if c.row in [1,4]:
            c.fill=PatternFill('solid',fgColor='17324D');c.font=Font(name='맑은 고딕',size=18 if c.row==1 else 11,bold=True,color='FFFFFF')
        elif c.row>=5 and ('확인' in str(c.value) or '예정' in str(c.value)):
            c.fill=PatternFill('solid',fgColor='FFF2D5')
for r in range(1,ws.max_row+1):ws.row_dimensions[r].height=44 if r>=5 else 34
table=Table(displayName='SimpleProjectList',ref=f'A4:D{ws.max_row}')
table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(table)
ws.freeze_panes='B5';ws.sheet_view.showGridLines=False;ws.sheet_view.zoomScale=85
ws.print_title_rows='1:4';ws.print_area=f'A1:D{ws.max_row}'
ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.sheet_properties.pageSetUpPr.fitToPage=True
for s in list(wb)[1:]:s.sheet_state='hidden'
wb.active=0;wb.save(path)
check=load_workbook(path)
assert check.active.max_column==4 and check.active.max_row==57
assert len([s for s in check if s.sheet_state=='visible'])==1
for name in ['회사실적_문서근거','대표경력_개인']:
    for r in check[name].iter_rows(min_row=5,values_only=True):assert r==source[r[0]]
font=ImageFont.truetype('C:/Windows/Fonts/malgun.ttf',16)
im=Image.new('RGB',(sum(widths)*7+20,570),'white');d=ImageDraw.Draw(im)
y=10
for r in range(1,10):
    x=10;h=58
    if r<4:
        d.text((x+8,y+12),str(check.active.cell(r,1).value),font=font,fill='#17324D')
    else:
        for c,w in enumerate(widths,1):
            w*=7;text=str(check.active.cell(r,c).value or '');lines=['']
            for ch in text:
                if font.getlength(lines[-1]+ch)>w-16:lines.append('')
                lines[-1]+=ch
            d.rectangle((x,y,x+w,y+h),fill='#17324D' if r==4 else '#EDF3F7' if r%2 else 'white')
            d.multiline_text((x+8,y+9),'\n'.join(lines),font=font,fill='white' if r==4 else '#233D50',spacing=3);x+=w
    y+=h
im.save(out/'simple-preview.png')
print('Updated existing file: 53 rows, 4 columns; 4 repeated personal entries linked in comments; sources retained.')
